"""Read an unmodified Google Sheets Excel download into migration JSON.

Usage: python scripts/convert-sheet-export.py migration-data/workbook.xlsx
Requires openpyxl for reading only. Never rewrites the original workbook.
"""
import datetime as dt
import hashlib
import json
import re
import sys
from pathlib import Path

IDENTIFIERS = {"id", "code", "storeCode", "sku", "barcode", "SKU", "Barcode", "Vendor Code", "Store Code", "Order ID"}
TABS = ("Products", "Vendors", "Orders", "Counter")
OLD_ORDER_HEADERS = ["Date", "Order ID", "Product", "SKU", "Barcode", "Displays", "Pairs", "Wholesale/Unit ($)", "Line Wholesale ($)", "SRP/Unit ($)", "Line SRP ($)"]
ORDER_HEADERS = ["Date", "Order ID", "Product", "SKU", "Barcode", "Order Unit", "Order Qty", "Total Units", "Dealer/Unit ($)", "Line Dealer ($)", "SRP/Unit ($)", "Line SRP ($)", "Order Sent", "Invoice Sent", "Payment Received", "Cancelled", "Company", "Vendor Code", "Store Code", "Customer Email", "Comments", "Agent Name", "Agent Email"]


def map_headers(name, headers):
    # Apps Script initializes headers only on empty sheets, but appends 23 fields.
    if name == "Orders" and headers == OLD_ORDER_HEADERS + [""] * 12:
        return ORDER_HEADERS.copy(), "legacy-headings-with-23-column-apps-script-layout"
    return headers, None


def read_value(cell, header, errors, sheet):
    value = cell.value
    if cell.data_type == "e":
        errors.append(f"{sheet}!{cell.coordinate}: spreadsheet error {value}")
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if header in IDENTIFIERS:
        if isinstance(value, bool):
            errors.append(f"{sheet}!{cell.coordinate}: boolean identifier")
        if isinstance(value, (int, float)):
            if not float(value).is_integer():
                errors.append(f"{sheet}!{cell.coordinate}: noninteger numeric identifier")
            if abs(value) >= 10**15:
                errors.append(f"{sheet}!{cell.coordinate}: numeric identifier may have lost precision; export as text")
            text = str(int(value)) if float(value).is_integer() else str(value)
            if re.fullmatch(r"0+", cell.number_format or ""):
                text = text.zfill(len(cell.number_format))
            return text
        return str(value)
    return value


def convert(path):
    import openpyxl
    values = openpyxl.load_workbook(path, read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
    errors = []
    result = {}
    mappings = {}
    try:
        missing = [name for name in TABS if name not in values.sheetnames]
        if missing:
            raise ValueError("Missing sheet tabs: " + ", ".join(missing))
        for name in TABS:
            ws = values[name]
            if ws.max_column is None:
                ws.calculate_dimension(force=True)
            if name == "Counter":
                result[name] = ws["A1"].value
                if result[name] is None:
                    errors.append("Counter!A1: missing order counter")
                continue
            rows = ws.iter_rows()
            header_cells = next(rows, [])
            headers = [str(c.value).strip() if c.value is not None else "" for c in header_cells]
            headers += [""] * max(0, ws.max_column - len(headers))
            headers, mapping = map_headers(name, headers)
            if mapping:
                mappings[name] = mapping
            nonempty = [h for h in headers if h]
            if not nonempty:
                errors.append(f"{name}: missing column headers")
            if len(nonempty) != len(set(nonempty)):
                errors.append(f"{name}: duplicate column headers")
            output = []
            for cells in rows:
                row = {}
                for i, cell in enumerate(cells):
                    header = headers[i] if i < len(headers) else ""
                    if not header:
                        if cell.value is not None:
                            errors.append(f"{name}!{cell.coordinate}: data under a blank header")
                        continue
                    row[header] = read_value(cell, header, errors, name)
                output.append(row)
            # Preserve interior blank rows so importer row references remain exact.
            while output and all(v == "" for v in output[-1].values()):
                output.pop()
            result[name] = output
        for name in TABS:
            for frow, vrow in zip(formulas[name].iter_rows(), values[name].iter_rows()):
                for fcell, vcell in zip(frow, vrow):
                    if fcell.data_type == "f" and vcell.value is None:
                        errors.append(f"{name}!{fcell.coordinate}: formula has no cached result; recalculate in Sheets and download again")
        if errors:
            raise ValueError("\n".join(errors))
        result["_source"] = {"filename": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "convertedAt": dt.datetime.now(dt.timezone.utc).isoformat(), "legacyTimeZone": "America/Toronto", "headerMappings": mappings}
        return result
    finally:
        values.close()
        formulas.close()


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("Usage: python scripts/convert-sheet-export.py migration-data/workbook.xlsx")
    source = Path(sys.argv[1])
    try:
        result = convert(source)
        destination = source.with_suffix(".json")
        with destination.open("x", encoding="utf-8") as stream:
            json.dump(result, stream, indent=2, ensure_ascii=False)
        print(json.dumps({"converted": str(destination), "rows": {name: len(result[name]) for name in TABS if name != "Counter"}, "originalWorkbookUnchanged": True}))
    except (ValueError, FileExistsError) as error:
        sys.exit(str(error))
